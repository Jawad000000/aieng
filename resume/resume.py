import os
from pathlib import Path
import time
from groq import Groq
from dotenv import load_dotenv
from pydantic import BaseModel
import json

load_dotenv()
my_api_key=os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("no api key")

client=Groq(api_key=my_api_key)
model="qwen/qwen3.8-27b"

job_description="""
Do you want to solve real customer problems through innovative technology? Do you enjoy working on scalable services in a collaborative team environment? Do you want to see your code directly impact millions of customers worldwide?

At Amazon, we hire the best minds in technology to innovate and build on behalf of our customers. Customer obsession is part of our company DNA, which has made us one of the world's most beloved brands.

Our Software Development Engineers (SDEs) use modern technology to solve complex problems while seeing their work's impact first-hand. The challenges SDEs solve at Amazon are meaningful and influence millions of customers, sellers, and products globally. We seek individuals passionate about creating new products, features, and services while managing ambiguity in an environment where development cycles are measured in weeks, not years.

At Amazon, we believe in ownership at every level. As an SDE-I, you'll own the entire lifecycle of your code - from design through deployment and ongoing operations. This ownership mindset, combined with our commitment to operational excellence, ensures we deliver the highest quality solutions for our customers.

We're looking for curious minds who think big and want to define tomorrow's technology. At Amazon, you'll grow into the high-impact engineer you know you can be, supported by a culture of learning and mentorship. Every day brings exciting new challenges and opportunities for personal growth.

We will take your location preferences into consideration. Preferences are based on business availability and are not guaranteed. Applicants will be considered at all locations in the United States including but not limited to:

• AZ (Phoenix, Tempe)
• CA (Berkeley, Culver City, Cupertino, East Palo Alto, Irvine, Los Angeles, Manhattan Beach, Palo Alto, San Diego, San Francisco, San Jose, San Luis Obispo, Santa Barbara, Santa Clara, Santa Cruz, Santa Monica, Sunnyvale)
• CO (Boulder, Denver)
• GA (Atlanta, Kennesaw)
• IL (Chicago)
• MA (Boston, Cambridge, Hudson, North Reading, Westborough)
• MD (Baltimore)
• MI (Detroit)
• MN (Minneapolis)
• NJ (Jersey City)
• NY (New York)
• OR (Portland)
• PA (Philadelphia, Pittsburgh)
• TN (Nashville)
• TX (Austin, Dallas)
• VA (Arlington, Herndon)
• WI (Madison)
• WA (Bellevue, Seattle, Redmond)
***Locations are subject to change***

During your application, you will have the opportunity to highlight your expertise in one or more of these specialized areas:

• Machine Learning • Distributed Systems and Data Management • Database Systems • Quantum Computing • Network Development • Query Processing and Optimization • Automated Reasoning • Embedded Systems • Data Engineering • Mobile Development • Game Development

While team placement is based on business needs, sharing your specific skills and interests helps us better align your experience with potential roles.

Applications are reviewed on a rolling basis. For an update on your status or to confirm your application was submitted successfully, please login to your candidate portal at Amazon.jobs. Amazon works with a high volume of applicants so we appreciate your patience as we review applications.


Key job responsibilities
• Collaborate and communicate effectively with experienced cross-disciplinary Amazonians to design, build, and operate innovative products and services that delight our customers, while participating in technical discussions to drive solutions forward.
• Design and develop scalable solutions using cloud-native architectures and microservices in a large distributed computing environment.
• Participate in code reviews and contribute to technical documentation.
• Build and maintain resilient distributed systems that are scalable, fault-tolerant, and cost-effective.
• Leverage and contribute to the development of GenAI and AI-powered tools to enhance development productivity while staying current with emerging technologies.
• Write clean, maintainable code following best practices and design patterns.
• Work in an agile environment practicing CI/CD principles while participating in operational responsibilities including on-call duties.
• Demonstrate operational excellence through monitoring, troubleshooting, and resolving production issues.
Basic Qualifications

- Experience with at least one general-purpose programming language such as Java, Python, C++, C#, Go, Rust, or TypeScript
- Experience with data structure implementation, basic algorithm development, and/or object-oriented design principles
- Currently has, or is in the process of obtaining a bachelor’s degree in Computer Science, Computer Engineering, Data Science, Information Systems, or related STEM fields
- Must be 18 years of age of older
Preferred Qualifications

- Experience from previous technical internship(s) or demonstrated project experience
- Experience with one or more of the following: AI tools for development productivity, Cloud platforms (preferably AWS), Database systems (SQL and NoSQL), Contributing to open-source projects, Version control systems, Debugging and troubleshooting complex systems
- Demonstrated ability to learn and adapt to new technologies quickly
- Basic understanding of software development lifecycle (SDLC)
- Strong problem-solving and analytical skills
- Excellent written and verbal communication skills
"""

class Job(BaseModel):
    role: str
    required_skills:list[str]
    preferred_skills:list[str]
    minimum_experience:float | None
    educational_requirements:list[str]
    responsibilities:list[str]

job_Schema = Job.model_json_schema()

system_prompt =f"""
you are an expert hr assistant

Your job is to analyze job description and extract
structured information from them.

Return only valid JSON matching this schema:
{job_Schema}
IMPORTANT:
do not return the schema itself.
Do not return fields like "properties, "title" or "type".
FILL THE SCHEMA WITH ACTUAL INFORMATION EXTRACTED FROM THE JOB DESCRIPTION.

if minimum experience is not written return null.
if any information for a list is missing, return a empty list
do not invent information
"""

user_prompt=f"""
Analyze the following job description
{job_description}
"""

message_system={
    "role":"system",
    "content": system_prompt

}

message_user={
    "role":"user",
    "content": user_prompt
}

response_format={
    "type" : "json_object"
}

message = [message_system, message_user]

response = client.chat.completions.create(model=model, messages=message, response_format=response_format)

answer = response.choices[0].message.content

raw_json = answer
job_data = json.loads(raw_json)
job = Job(**job_data)

print(job.minimum_experience)
print(job.educational_requirements)

class MatchDetails(BaseModel):
    candidate_name: str
    matching_skills: list[str]
    missing_skills: list[str]
    experience_requirements_met: bool
    final_verdict: str

class MatchResults(BaseModel):
    score: float
    details: MatchDetails

class Experience(BaseModel):
    company:str | None =None
    role:str | None =None
    duration:float | None =None
    description:str | None =None
    skills_used:list[str] =[]

class Resume(BaseModel):
    name:str | None =None
    email:str | None =None
    phone:str | None =None
    education:list[str] =[]
    experiences:list[Experience] =[]
    skills:list[str] =[]
    total_years_of_experience:float | None =None
    projects:list[str] =[]
    certifications:list[str] =[]

resume_schema = Resume.model_json_schema()

def final_score(job, resume):
    prompt = f"""
    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    """
    
    messages = [
        {"role": "system", "content": "You are an expert HR evaluation assistant."},
        {"role": "user", "content": prompt}
    ]
    
    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format={
            "type": "json_schema",
            "json_schema": {
                "name": "candidate_matching",
                "schema": MatchResults.model_json_schema()
            }
        },
        max_tokens=600,
        temperature=0.1
    )
    
    return MatchResults.model_validate_json(response.choices[0].message.content)

def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.
    Extract structured information from the resume text.
    Return only valid JSON matching this schema:
    {resume_schema}
    IMPORTANT:
    Do not return the schema itself.
    Do not return fields like "properties", "title" or "type".
    FILL THE SCHEMA WITH ACTUAL INFORMATION EXTRACTED FROM THE RESUME TEXT.
    If any information for a list is missing, return an empty list.
    Do not invent information.
    """
    
    user_prompt = f"""
    parse the following resume text:
    {resume_text}
    """

    message_system = {
        "role": "system",
        "content": system_prompt
    }

    message_user = {
        "role": "user",
        "content": user_prompt
    }

    messages = [message_system, message_user]

    response_format = {
        "type": "json_object"
    }

    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    
    answer = response.choices[0].message.content
    raw_json = answer
    resume_data = json.loads(raw_json)
    
    return Resume(**resume_data)


#resume parse
from pypdf import PdfReader
from docx import Document

def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        text += page.extract_text()
    return text

def read_docx(file_path):
    doc = Document(file_path)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += cell.text + "\t"
            text += "\n"

    return text

def read_resume(file_path):
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return read_pdf(file_path)
    elif ext == ".docx":
        return read_docx(file_path)
    else:
        raise ValueError("Unsupported file format. Please provide a PDF or DOCX file.")
    
resume_folder = Path("resumes")
all_results = []

for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue
    print(f"Processing resume: {file_path.name}")
    resume_text = read_resume(file_path)
    parsed_resume = parse_resume(resume_text)
    time.sleep(5) 
    result = final_score(job, parsed_resume)
    time.sleep(5)
    print(f"Match score for {file_path.name}: {result.score}")
    all_results.append({
        "resume_file": file_path.name,
        "match_score": result.score,
        "details": result.details.model_dump()
    })
all_results.sort(key=lambda x: x["match_score"], reverse=True)

top2 = all_results[:2]
worst2 = all_results[-2:]

print("\nTop 2 candidates:")
for result in top2:
    print(f" - {result['resume_file']}: {result['match_score']}")

print("\nWorst 2 candidates:")
for result in worst2:
    print(f" - {result['resume_file']}: {result['match_score']}")